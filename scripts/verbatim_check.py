#!/usr/bin/env python3
"""
verbatim_check.py

Counts the verbatim word runs a chapter shares with its sources: the count
behind verification step 8 of the source-verifier, rated by the thresholds
of checklists/plagiarism_protection.md:

    5-7 words    🟡 polish
    8-14 words   ⚠️ plagiarism risk, must be reformulated
    15+ words    ❌ serious

Words are counted as written, each a stretch between whitespace or dashes
(–, —, a typed -- or ---): a hyphenated compound (long-term) or a word with an
apostrophe (city’s) is one word, although its parts are matched one by one,
and a dash between words separates them. Runs of 1-4 words are not
listed. A run counts across sentence boundaries and across citations, as
Pandoc citekeys ([@key, p. 5]) or in author-year form ((Keller 2018, S. 44),
(ebd.)); letter case and punctuation are ignored. A quotation of five or
more words in quotation marks („…“, “…”, "…", ‚…‘, ‘…’, '…', »…«, «…»),
Markdown block quotes and Word paragraphs in a quote style (also Pandoc's
"Block Text") are marked direct quotes and are not counted; the words of a
shorter quotation (a term in quotation marks) keep a run going but are not
counted in its length. Each direct quote (its parts between […] or (…)
omissions) is looked up in the sources instead; one that is not found
verbatim is listed for a check by hand, without changing the exit code. The
chapter's bibliography (from a heading such as "Literaturverzeichnis" or
"References" to the next heading of its level) is not counted if every entry
in it has a year, as in a reference list; a body section of that name, an
appendix and footnotes count.

Sources are files or folders, searched recursively; folders named
_guidelines (guidelines are not sources) and verification (the saved
reports) inside them are skipped:
  - extract files (.md): the quotes of the research-agent's extracts (the
    "Original" field, quoted text in the "Secondary citation?" field, and
    "- p. 42: ..." lines), with their page and quote ID. A Markdown table
    with an "Original quote" column, such as the source registry in the
    workspace/state.md of a seminar paper, is read row by row.
  - PDFs (.pdf): the text layer, read with pypdf, without the running
    headers and footers (also those alternating between left and right
    pages) and page numbers, so that a run goes on across a page break; a page here is the PDF viewer's page index, which may differ
    from the printed page number.
    A PDF without a text layer (a scan) and a ZIP archive named .pdf are
    not checked: they are covered only through the quotes in their extract.
  - plain-text files (.txt)
What cannot be read this way (a missing path, a Word, EPUB, HTML or image
file, the pages without text of a PDF that is half or more scans, PDFs
while pypdf is not installed) is listed under "Not checked"; other files,
hidden files and an .md file in a folder that is no extract (a handover or
README) are passed over.

The report's header gives the chapter's word count for the page budget, as
written: the stretches between whitespace or dashes that hold a letter or
digit (a hyphenated compound is one word), headings and block quotes
included. In Markdown, footnotes count too; the YAML front matter, HTML
comments, raw blocks for the export (```{=openxml}; the text of their w:t
elements is still compared with the sources), Pandoc citekeys (the
whole bracket, [@key, p. 5]: Pandoc writes the citation only at the
export), footnote marks and the bibliography do not (the line says when
it recognized none), while author-year citations written into the text
do. A Word chapter counts its body paragraphs, without tables, footnotes,
text boxes, content controls and the bibliography. --words prints only
that line and needs no sources.

With --chapters and the paths after it, the chapter is also compared with
the other chapters, for the self-plagiarism check between chapters (A.2 of
checklists/consistency.md): chapter files, or the .md and .docx files
directly in a folder, the chapter itself left out. Each is read as the
chapter is, so that its quotations, block quotes, bibliography and raw
blocks for the export are left out on its side too, and a run is found and
counted as above. Under "Runs shared with other chapters" the report lists
each run of 5 or more words the chapter shares with one of them, with that
chapter's file and line, unless the run lies wholly within a run from the
sources listed above; and each direct quote of the chapter that the sources
do not have but another chapter has outside quotation marks (the research
question quoted again, for example). These are not rated and do not change
the exit code. Another chapter that cannot be read is listed under "Could
not be read", as a source is; a path that does not exist, under "Not
checked".

Usage:
    python scripts/verbatim_check.py <chapter> <source> [<source> ...] [--chapters <chapter or folder> ...]
    python scripts/verbatim_check.py --words <chapter>
    python scripts/verbatim_check.py workspace/output/chapters/ch_2_basics.md workspace/sources/
    python scripts/verbatim_check.py workspace/output/chapters/ch_3_conclusion.md workspace/sources/ --chapters workspace/output/chapters/
    python scripts/verbatim_check.py --help

The chapter may be Markdown (.md), plain text (.txt) or Word (.docx).

Exit codes: 0 = no run of 8 or more words from the sources, 1 = at least
one run of 8 or more words from the sources (⚠️ or ❌), 2 = bad arguments, a
chapter that cannot be read or holds no text, a source or another chapter
that cannot be read (the report is printed all the same, with the file
under "Could not be read"), or no source text to compare against. With
--words: 0, or 2 for bad arguments or a chapter that cannot be read.
"""

import bisect
import contextlib
import html
import io
import os
import re
import sys
import unicodedata
from pathlib import Path
from typing import NamedTuple, Optional

# Encoding guard for Windows consoles (CP1252 by default); a file name that is
# not valid UTF-8 is printed with escapes instead of ending in a traceback.
for stream in (sys.stdout, sys.stderr):
    try:
        if stream.encoding and stream.encoding.lower() != "utf-8":
            stream.reconfigure(encoding="utf-8", errors="backslashreplace")
        else:
            stream.reconfigure(errors="backslashreplace")
    except Exception:
        pass

# Thresholds of checklists/plagiarism_protection.md, section 1.
POLISH, RISK, SERIOUS = 5, 8, 15

USAGE = ("Usage: python scripts/verbatim_check.py <chapter> <source> [<source> ...] [--chapters <chapter or folder> ...]\n"
         "       python scripts/verbatim_check.py --words <chapter>")
CHAPTER_SUFFIXES = {".md", ".markdown", ".txt", ".docx"}
# The chapter files of a folder given to --chapters, as state_check.py reads one
FOLDER_SUFFIXES = {".md", ".markdown", ".docx"}
SOURCE_SUFFIXES = {".md", ".markdown", ".txt", ".pdf"}
# Document types a source might come in that the script cannot read: listed as
# not checked. Other files in a source folder (.bib, a saved page's scripts and
# fonts, shortcuts) are passed over.
LISTED_SUFFIXES = {".doc", ".docx", ".odt", ".rtf", ".pages", ".epub", ".mobi", ".azw3", ".djvu",
                   ".xps", ".zip", ".html", ".htm", ".mhtml", ".tif", ".tiff", ".jpg", ".jpeg", ".png"}
# Guidelines are not sources; verification/ holds the saved reports, which quote the
# chapter; __MACOSX/ holds a Mac archive's resource forks.
SKIPPED_FOLDERS = {"_guidelines", "verification", "__MACOSX"}

WORD = re.compile(r"[^\W_]+")
# A quotation within one paragraph: in double marks, German „…“ (also „…”, and
# „…" when no “ or ” follows), English “…”, straight "…", and guillemets in both
# directions; in single marks, German ‚…‘ (also ‚…’), English ‘…’, straight '…'
# and single guillemets. A ’ or ' followed by a letter is an apostrophe (city’s,
# don't), not a closing mark; a straight ' opens a quotation only after a space,
# a bracket or a dash, and neither ' nor ‘ opens one before a digit ('90s).
QUOTED = re.compile(r'„[^„“”]*[“”]|„[^„“”"]*"|“[^“”„]*”|"[^"]*"|»[^»«]*«|«[^«»]*»'
                    r"|‚(?:[^‚‘’]|’(?=[^\W_]))*[‘’]|‘(?!\d)(?:[^‘’]|’(?=[^\W_]))*’"
                    r"|(?<![^\s(\[—–-])'(?!\d)(?:[^']|'(?=[^\W_]))*'(?![^\W_])|›[^›‹]*‹|‹[^‹›]*›")
# Omissions and insertions inside a quotation split it into parts that are
# looked up separately: […], [...], (…), (...), …, ..., and [any insertion].
OMISSION = re.compile(r"\[[^\]]*\]|\((?:…|\.{3})\)|…|\.{3}")
# A bracketed letter that adapts a quotation to the sentence ("[d]ie Wirkung") is
# part of its word, not an omission.
ADAPTED_LETTER = re.compile(r"(?<![^\W_])\[([^\W\d_])\](?=[^\W_])")
PANDOC_CITATION = re.compile(r"\[[^\[\]]*@[^\[\]]*\]")
# A citation in parentheses, author-year style: "(Keller 2018, S. 44)", "(vgl. Keller
# & Weber 2018; Smith 2020, p. 5)", "(2018)", "(ebd., S. 45)", "(Ebd.: 45)". Like a
# citekey it is skipped, in the chapter and in the sources, so that a run continues
# across it.
CITATION = re.compile(
    r"\((?:(?:vgl|cf|see|siehe)\.?\s+)?"
    r"(?:[A-ZÄÖÜ][^()]{0,150}?(?:\b\d{4}[a-z]?\b|\bo\.\s?J\.|\bn\.\s?d\.)[^()]{0,80}"
    r"|\d{4}[a-z]?(?:\s*[,:;][^()]{0,80})?"
    r"|(?:[Ee]bd|[Ee]benda|[Ii]bid(?:em)?)\.?(?:[,:\s][^()]{0,80})?)\)")
HTML_COMMENT = re.compile(r"<!--.*?-->", re.S)
# A raw block for one output format (```{=openxml}, also with ~~~ and a longer
# fence): markup that Pandoc passes to the export, such as the bookmark that
# starts the back matter (docs/reference_management.md), not text. Like an HTML
# comment it is blanked, so it neither counts nor joins a run. It closes with a
# fence of the same character, at least as long.
# The text of a Word raw block (the w:t elements of ```{=openxml}; Pandoc drops
# a {=docx} block): the export writes it into the document, so it is compared
# with the sources, though not counted for the page budget. A paragraph ends at
# </w:p> or is an empty <w:p/>; in it, a tab or line break shows as a space and
# a non-breaking hyphen as a hyphen. A text box's paragraphs are paragraphs of
# their own, and the copy of it for older readers (mc:Fallback) is left out.
# Pandoc's block-quote style (BlockText, the one quote style of its reference
# document and of scripts/build_reference_docx.py without --base) makes a
# paragraph a quote; a style the reference document lacks is plain text in the
# export, and one a --base template adds (Zitat) is compared as plain text too. Attribute
# scans stop at the next tag, so that malformed markup takes no extra time.
WORD_PARAGRAPH_END = re.compile(r"</w:p>|<w:p(?:\s[^<>]*)?/>")
WORD_RUN_TEXT = re.compile(r"<w:t(?:\s[^<>]*)?>([^<]*)</w:t>"
                           r"|<w:(tab|ptab|br|cr|noBreakHyphen)\b[^<>]*(?:/>|>\s*</w:\2>)")
WORD_STYLE = re.compile(r"<w:pStyle\s[^<>]*?w:val=\"([^\"<>]*)\"")
RAW_QUOTE_STYLE = re.compile(r"(?i)(?:footnote)?blocktext")
QUOTE_STYLE = re.compile(r"(?i)quot|zitat|block ?text")
RAW_FORMAT = re.compile(r"\{=([^{}\s]+)\}")
RAW_BLOCK = re.compile(r"(?ms)^ {0,3}(?P<fence>(?P<char>[`~])(?P=char){2,})[ \t]*\{=[^{}\s]+\}[ \t]*\n"
                       r".*?^ {0,3}(?P=fence)(?P=char)*[ \t]*$")
# The heading of a chapter's bibliography (as in bibliography_check.py, and more):
# the titles listed under it are no runs.
BIBLIOGRAPHY_HEADINGS = {"bibliography", "references", "reference list", "works cited", "literature",
                         "sources", "literaturverzeichnis", "literatur", "verwendete literatur",
                         "bibliographie", "quellenverzeichnis", "quellen", "quellen- und literaturverzeichnis",
                         "literatur- und quellenverzeichnis"}
# A year (or "no date") in a paragraph: every entry of a reference list has one.
YEAR = re.compile(r"\b(?:1[5-9]|20)\d{2}[a-z]?\b|\bo\.\s?J\.|\bn\.\s?d\.", re.I)
ROMAN = r"(?:m{0,3}(?:cm|cd|d?c{0,3})(?:xc|xl|l?x{0,3})(?:ix|iv|v?i{0,3}))"
# A line of a PDF page that is only its number: "117", "- 12 -", "Seite 3", "iv", "3 / 10"
PAGE_NUMBER_LINE = re.compile(r"(?i)\s*[-–—]?\s*(?:(?:seite|page|s\.|p\.)\s*)?(?:\d{1,4}|(?=[ivxlcdm])" + ROMAN + r")"
                              r"(?:\s*(?:/|of|von)\s*\d{1,4})?\s*[-–—]?\s*")
FOOTNOTE_MARK = re.compile(r"\[\^[^\]]*\]")
# Markup that the word count leaves out: a heading's attributes ({#sec-intro
# .unnumbered}) and the fence lines of a code block or a Pandoc div (::: {...}).
MARKUP = re.compile(r"(?m)^(\s{0,3}#{1,6}\s.*?)\s*\{[^{}]*\}\s*$|^\s{0,3}(?:`{3,}|~{3,}|:{3,}).*$")


def normalize(text: str) -> str:
    """NFKC (ligatures such as ﬁ become fi), soft hyphens dropped."""
    return unicodedata.normalize("NFKC", text).replace("­", "")


def tokens(text: str) -> list:
    return [m.group(0).casefold() for m in WORD.finditer(text)]


# A dash between words (figure, en and em dash, horizontal bar, or a typed -- or
# ---) separates them like a space; a single hyphen joins a compound.
DASH = re.compile(r"[\u2012-\u2015]|-{2,}")


def spaced(text: str) -> str:
    """The text with each dash replaced by spaces of its length, so that offsets
    stay valid. For counting words as written only, never for matching: tokens()
    splits at a dash already."""
    return DASH.sub(lambda m: " " * len(m.group(0)), text)


def as_written(text: str) -> int:
    """The words as written: the stretches between whitespace or dashes that
    hold a letter or digit, so that a hyphenated compound (long-term) or a word
    with an apostrophe (city's) is one word, though tokens() splits it."""
    return sum(1 for w in spaced(text).split() if WORD.search(w))


def quote_parts(quote: str) -> list:
    """The parts of a quotation between omissions, as text."""
    return OMISSION.split(ADAPTED_LETTER.sub(r"\1", quote))


def read_file_text(path: Path) -> str:
    """UTF-8 (with or without BOM), UTF-16 with BOM, otherwise Windows-1252."""
    data = path.read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def one_line(exc: Exception) -> str:
    return " ".join(str(exc).split())


def fail(message: str) -> None:
    print(f"Error: {message}", file=sys.stderr)
    sys.exit(2)


# ---------------------------------------------------------------------------
# Chapter
# ---------------------------------------------------------------------------

class Chapter(NamedTuple):
    words: list            # casefolded words, None where a run must break
    places: list           # for each entry of words: (line or paragraph number) or None
    spans: list            # for each entry of words: (start, end) in text, or None
    text: str              # the normalized text the spans refer to
    unit: str              # "line" or "paragraph"
    quotes: list           # (place, quotation text) of the marked direct quotes
    quoted: list           # for each entry of words: True inside a short quotation
    bibliography: Optional[tuple] = None  # its first and last line or paragraph (None: to the end)
    counted: Optional[int] = None  # text[:counted] is what the word count counts (None: all of it)


def blank(text: str, start: int, end: int) -> str:
    """Replace text[start:end] by spaces, keeping line breaks (and line numbers)."""
    return text[:start] + re.sub(r"[^\n]", " ", text[start:end]) + text[end:]


def bibliography_heading(text: str) -> bool:
    """
    A heading such as "Literaturverzeichnis", "7 References", "VII. Literatur"
    or "**Bibliography:**". One numbered as a subsection ("2.1 Literature")
    belongs to the body.
    """
    key = text.strip().strip("*_ ").strip()
    if re.match(r"\d+\.\d", key):
        return False
    key = re.sub(r"^(?:\d+|" + ROMAN + r"|[a-z])\.?\s+", "", key, count=1, flags=re.I)
    return key.lower().rstrip(".:").strip() in BIBLIOGRAPHY_HEADINGS


def reference_list(paragraphs: list) -> bool:
    """Every paragraph of the section has a year, apart from short sub-headings."""
    paragraphs = [p for p in paragraphs if p.strip()]
    return any(YEAR.search(p) for p in paragraphs) and all(
        YEAR.search(p) or len(p.split()) <= 6 for p in paragraphs)


def bibliography_section(headings: list, count: int) -> Optional[tuple]:
    """
    The (start, end) of a chapter's bibliography among its items 0..count-1,
    end exclusive. headings: (index, level, rank, text) of its headings, level
    1 the highest. As in bibliography_check.py, the highest-level bibliography
    heading wins (rank breaks a tie, then the order), and the section ends at
    the next heading of the same or a higher level that is not a bibliography
    heading, so that an appendix after it still counts.
    """
    found = [h for h in headings if bibliography_heading(h[3])]
    if not found:
        return None
    start, level, _, _ = min(found, key=lambda h: (h[1], h[2], h[0]))
    end = next((i for i, lv, _, text in headings
                if i > start and lv <= level and not bibliography_heading(text)), count)
    return start, end


def element_spans(block: str, name: str) -> list:
    """
    (start, end) of each outermost element called name in block, in one pass:
    an element inside another of its kind belongs to the outer one, an empty
    one (<name/>) is a span of its own tag, and one without its end tag runs to
    the end of the block.
    """
    out, depth, start = [], 0, 0
    for tag in re.finditer(r"<(/?)" + re.escape(name) + r"\b[^<>]*?(/?)>", block):
        if tag.group(1):
            if depth:
                depth -= 1
                if not depth:
                    out.append((start, tag.end()))
        elif tag.group(2):
            if not depth:
                out.append((tag.start(), tag.end()))
        else:
            if not depth:
                start = tag.start()
            depth += 1
    if depth:
        out.append((start, len(block)))
    return out


def blank_spans(text: str, spans: list) -> str:
    """text with each of the sorted, separate spans replaced as blank() does, in one pass."""
    parts, pos = [], 0
    for a, b in spans:
        parts += [text[pos:a], re.sub(r"[^\n]", " ", text[a:b])]
        pos = b
    return "".join(parts) + text[pos:]


def raw_paragraphs(block: str) -> list:
    """
    (text, offset of its first text, is_quote) of each paragraph of a Word raw
    block. A text box's paragraphs (w:txbxContent) are taken on their own, so
    that the paragraph around it keeps its text before and after the box and
    its own style; the copy of a box for older readers (mc:Fallback) is left out.
    """
    block = blank_spans(block, element_spans(block, "mc:Fallback"))
    spans = element_spans(block, "w:txbxContent")
    boxes = [(a, block[a:b]) for a, b in spans]
    block = blank_spans(block, spans)
    out = []
    for base, part_of in [(0, block)] + boxes:
        start = 0
        for end in [e.start() for e in WORD_PARAGRAPH_END.finditer(part_of)] + [len(part_of)]:
            part = part_of[start:end]
            chars, at = [], None
            for r in WORD_RUN_TEXT.finditer(part):
                if r.group(1) is not None:
                    at = r.start() if at is None else at
                    chars.append(html.unescape(r.group(1)))
                else:
                    chars.append("-" if r.group(2) == "noBreakHyphen" else " ")
            words = "".join(chars).strip()
            if words and at is not None:
                style = WORD_STYLE.search(part)
                out.append((words, base + start + at, bool(style and RAW_QUOTE_STYLE.fullmatch(style.group(1)))))
            start = end
    return out


def read_markdown(path: Path) -> tuple:
    """
    Paragraph-sized pieces of a Markdown or text chapter, (text, first line,
    is_quote), and the first and last line of its bibliography, if it has one.
    """
    text = normalize(read_file_text(path)).replace("\r\n", "\n").replace("\r", "\n")
    # YAML front matter at the top
    m = re.match(r"---\n.*?\n(?:---|\.\.\.)\n", text, re.S)
    if m:
        text = blank(text, 0, m.end())
    # footnote definitions ("[^1]: ..." and their indented lines) are body text
    footnotes, inside = set(), False
    for n, line in enumerate(text.split("\n")):
        inside = bool(re.match(r"\[\^[^\]]+\]:", line)) or (inside and bool(re.match(r"\s{2,}\S|\t", line)))
        if inside:
            footnotes.add(n)
    # the text of Word raw blocks that the export writes: one piece per w:p
    # paragraph, at the line of its first text, its runs joined as Word shows
    # them (a tab or line break between runs as a space); a paragraph in
    # Pandoc's block-quote style is a quote, as in the export; not in an HTML
    # comment, which the export drops
    comments = [(m.start(), m.end()) for m in HTML_COMMENT.finditer(text)]
    breaks = [m.start() for m in re.finditer("\n", text)]
    raw_text = []
    for m in RAW_BLOCK.finditer(text):
        fmt = RAW_FORMAT.search(m.group(0))
        if not fmt or fmt.group(1).lower() != "openxml":
            continue
        if any(a <= m.start() < b for a, b in comments):
            continue
        for words, at, is_quote in raw_paragraphs(m.group(0)):
            line = bisect.bisect_left(breaks, m.start() + at) + 1
            raw_text.append((normalize(words), line, is_quote, True))
    for pattern in (RAW_BLOCK, HTML_COMMENT, PANDOC_CITATION, FOOTNOTE_MARK):
        for m in reversed(list(pattern.finditer(text))):
            text = blank(text, m.start(), m.end())
    pieces = []
    lines = text.split("\n")
    # The bibliography is a list of titles: its lines are left out (blanked, so
    # that the line numbers stay) if it looks like one.
    headings = []
    for n, line in enumerate(lines):
        h = re.match(r"\s{0,3}(#{1,6})\s+(.*?)\s*(?:\{[^}]*\})?\s*#*\s*$", line)
        if h:
            headings.append((n, len(h.group(1)), 0, h.group(2)))
    cut = bibliography_section(headings, len(lines))
    if cut:
        body = [n for n in range(cut[0] + 1, cut[1]) if n not in footnotes]
        paragraphs = "\n".join(lines[n] if n in body else "" for n in range(cut[0] + 1, cut[1])).split("\n\n")
        # a raw paragraph in the section is one of its entries, as in the export
        paragraphs += [r[0] for r in raw_text if cut[0] + 1 <= r[1] <= cut[1]]
        if reference_list(paragraphs):
            for n in [cut[0]] + body:
                lines[n] = ""
        else:
            cut = None
    buf, first, quote = [], 0, False
    for n, line in enumerate(lines, 1):
        # a line without ">" right after a block-quote line continues the quote
        # (lazy continuation), unless it starts a heading, list item or fence
        is_quote = bool(re.match(r"\s{0,3}>", line)) or (
            quote and bool(buf) and bool(line.strip())
            and not re.match(r"\s{0,3}(?:#|[-*+]\s|\d+[.)]\s|```|~~~)", line))
        if not line.strip() or (buf and is_quote != quote):
            if buf:
                pieces.append(("\n".join(buf), first, quote))
            buf = []
            if not line.strip():
                continue
        if not buf:
            first, quote = n, is_quote
        buf.append(re.sub(r"^\s{0,3}>\s?", "", line) if is_quote else line)
    if buf:
        pieces.append(("\n".join(buf), first, quote))
    # 1-based first and last line, the last None when it runs to the end; the
    # text of Word raw blocks comes last, marked as not counted, and not from
    # the bibliography
    if cut:
        raw_text = [r for r in raw_text if not cut[0] + 1 <= r[1] <= cut[1]]
    return pieces + raw_text, (cut and (cut[0] + 1, cut[1] if cut[1] < len(lines) else None))


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC_NS = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
# What a paragraph's text leaves out: deleted and moved-away text of tracked
# changes, a text box's paragraphs and their copy for older readers (text boxes
# are not counted), the tab stops of its paragraph properties, and the small
# annotation above ruby text (its base is the line's text).
DOCX_SKIP = {W_NS + "del", W_NS + "moveFrom", W_NS + "txbxContent", MC_NS + "Fallback", W_NS + "pPr",
             W_NS + "rt"}


def docx_text(p) -> str:
    """
    A Word paragraph's text (its w:p element) as Word shows it with its
    tracked changes accepted: inserted text counts and deleted text does not; a
    tab is a tab, a line, column or page break a line break, a symbol from a
    symbol font (Insert > Symbol) a space, and a non-breaking hyphen a hyphen.
    """
    out = []

    def walk(el):
        for child in el:
            tag = child.tag
            if not isinstance(tag, str) or tag in DOCX_SKIP:
                continue
            if tag == W_NS + "t":
                out.append(child.text or "")
            elif tag in (W_NS + "tab", W_NS + "ptab"):
                out.append("\t")
            elif tag in (W_NS + "br", W_NS + "cr"):
                out.append("\n")
            elif tag == W_NS + "noBreakHyphen":
                out.append("-")
            elif tag == W_NS + "sym":
                out.append(" ")
            else:
                walk(child)
    walk(p)
    return "".join(out)


def read_docx(path: Path) -> tuple:
    """
    Paragraphs of a Word chapter, (text, paragraph number, is_quote), and the
    first and last paragraph of its bibliography, if it has one.
    """
    try:
        from docx import Document
    except ImportError:
        fail("python-docx not installed. Install: pip install -r requirements.txt")
    try:
        doc = Document(str(path))
    except Exception as exc:
        fail(f"cannot read {path} as .docx ({type(exc).__name__}: {one_line(exc)})")
    pieces, headings = [], []
    carry, carry_n, carry_q = "", None, False
    try:
        for n, p in enumerate(doc.paragraphs, 1):
            style = (p.style.name if p.style is not None else "") or ""
            # Quote, Intense Quote, Zitat; Pandoc's "Block Text"; LibreOffice's
            # "Quotations" / "Block Quotation"
            is_quote = bool(QUOTE_STYLE.search(style))
            text = docx_text(p._p)
            if carry_n is not None:
                # numbered where the joined text starts (a wholly deleted paragraph gives none)
                text, n, carry, carry_n = carry + text, (carry_n if carry.strip() else n), "", None
            # a paragraph mark deleted with Track Changes on joins the paragraph to the
            # next one once the changes are accepted, in the next one's style
            mark = p._p.find(W_NS + "pPr/" + W_NS + "rPr")
            if mark is not None and (mark.find(W_NS + "del") is not None or mark.find(W_NS + "moveFrom") is not None):
                carry, carry_n, carry_q = text, n, is_quote
                continue
            if text.strip():
                pieces.append((normalize(text), n, is_quote))
                # heading level from the style ("Heading 2", "Überschrift 2", "Title");
                # a plain paragraph "Literatur" counts as level 1 but yields to a styled one
                h = re.match(r"(?i)(?:heading|überschrift)\s*(\d)", style)
                if h or re.match(r"(?i)title|titel", style):
                    headings.append((n, int(h.group(1)) if h else 0, 0, normalize(text)))
                elif bibliography_heading(text):
                    headings.append((n, 1, 1, normalize(text)))
        if carry_n is not None and carry.strip():       # the last paragraph's mark: nothing to join
            pieces.append((normalize(carry), carry_n, carry_q))
    except Exception as exc:
        fail(f"cannot read {path} as .docx ({type(exc).__name__}: {one_line(exc)})")
    # The bibliography is a list of titles: left out if it looks like one.
    last = len(doc.paragraphs) + 1
    cut = bibliography_section(headings, last)
    if cut and reference_list([text for text, n, _ in pieces if cut[0] < n < cut[1]]):
        pieces = [piece for piece in pieces if not cut[0] <= piece[1] < cut[1]]
        return pieces, (cut[0], cut[1] - 1 if cut[1] < last else None)
    return pieces, None


def read_chapter(path: Path) -> Chapter:
    docx = path.suffix.lower() == ".docx"
    pieces, cut = read_docx(path) if docx else read_markdown(path)
    words, places, spans, quotes, inside = [], [], [], [], []
    text_parts, offset, counted = [], 0, None

    def brk():
        if words and words[-1] is not None:
            words.append(None)
            places.append(None)
            spans.append(None)
            inside.append(False)

    for piece, first, is_quote, *raw in pieces:
        if raw and counted is None:
            counted = offset             # from here on: the text of raw blocks, not counted
        base = offset
        text_parts.append(piece + "\n\n")
        offset += len(piece) + 2
        if is_quote:
            quotes.append((first, piece))
            brk()
            continue
        # A quotation of fewer than POLISH words (a short quote, a technical term,
        # scare quotes) does not break a run, which would cut a copied sentence
        # around it into pieces; its words are not counted in the run's length.
        quoted, short = [], []
        for m in QUOTED.finditer(piece):
            (quoted if as_written(m.group(0)) >= POLISH else short).append((m.start(), m.end()))
        for s, e in quoted:
            quotes.append((first + (0 if docx else piece.count("\n", 0, s)), piece[s + 1:e - 1]))
        # a citation is skipped without breaking the run
        cited = [(m.start(), m.end()) for pattern in (CITATION, PANDOC_CITATION)
                 for m in pattern.finditer(piece)]
        pos = 0
        for s, e in quoted + [(len(piece), len(piece))]:
            for m in WORD.finditer(piece, pos, s):
                if any(a <= m.start() < b for a, b in cited):
                    continue
                words.append(m.group(0).casefold())
                places.append(first + (0 if docx else piece.count("\n", 0, m.start())))
                spans.append((base + m.start(), base + m.end()))
                inside.append(any(a < m.start() < b for a, b in short))
            brk()
            pos = e
    return Chapter(words, places, spans, "".join(text_parts), "paragraph" if docx else "line",
                   quotes, inside, cut, counted)


def word_count(chapter: Chapter) -> str:
    """
    The chapter's words as written, for the page budget: the stretches between
    whitespace or dashes that hold a letter or digit (long-term, city's: one
    word each) in what read_chapter keeps, headings, block quotes and
    quotations included. The line names the bibliography as not counted only
    if read_chapter recognized one, and otherwise says that a reference list
    in the file counts.
    """
    text = chapter.text if chapter.counted is None else chapter.text[:chapter.counted]
    n = as_written(MARKUP.sub(lambda m: m.group(1) or "", text))
    if chapter.unit == "paragraph":
        kept, left = "headings and block quotes", ["tables", "footnotes", "text boxes", "content controls"]
    else:
        kept, left = ("headings, block quotes and footnotes",
                      ["YAML header", "HTML comments", "raw blocks", "citekeys", "footnote marks"])
    if chapter.bibliography:
        left.append("the bibliography")
    what = f"{kept} included; {', '.join(left[:-1])} and {left[-1]} not counted"
    if not chapter.bibliography:
        # none, or one under a heading it does not know or with an entry without a year
        what += "; no bibliography recognized: a reference list in the file counts"
    return f"Words: {n} (as written, {what})"


def checked_chapter(path: Path) -> Path:
    """The chapter argument, or exit 2 if it cannot be opened or is no chapter."""
    try:
        with open(path, "rb"):
            pass
    except OSError as exc:
        reason = "is a directory" if os.path.isdir(path) else exc.strerror or exc
        fail(f"cannot open {path}: {reason}")
    if path.suffix.lower() not in CHAPTER_SUFFIXES:
        fail(f"expected a .md, .txt or .docx chapter, got "
             f"{path.suffix or 'a file without extension'}: {path}")
    return path


def other_chapter_files(args: list, chapter_path: Path, unchecked: list) -> list:
    """
    The files of --chapters: each chapter file named, and the .md and .docx
    files directly in each folder named (not those in subfolders, no hidden
    file), without the checked chapter and without repeats; what is not a
    chapter goes to unchecked.
    """
    found, seen = [], {chapter_path.resolve()}
    for arg in args:
        p = Path(arg)
        if p.is_dir():
            try:
                files = sorted(x for x in p.iterdir() if x.is_file() and not x.name.startswith((".", "~$"))
                               and x.suffix.lower() in FOLDER_SUFFIXES)
            except OSError as exc:
                unchecked.append(f"{p} (folder not read: {exc.strerror or exc})")
                continue
        elif p.is_file():
            if p.suffix.lower() not in CHAPTER_SUFFIXES:
                unchecked.append(f"{p} ({p.suffix or 'a file without extension'} is not read: "
                                 "give a .md, .txt or .docx chapter)")
                continue
            files = [p]
        else:
            unchecked.append(f"{p} (not found)")
            continue
        for f in files:
            if f.resolve() not in seen:
                seen.add(f.resolve())
                found.append(f)
    return found


def read_other_chapter(path: Path) -> tuple:
    """(Chapter, None) for another chapter, or (None, why) if it cannot be read:
    read_chapter() ends the script with exit code 2, as it must for the checked
    chapter."""
    err = io.StringIO()
    try:
        with contextlib.redirect_stderr(err):
            return read_chapter(path), None
    except SystemExit:
        why = " ".join(err.getvalue().split()).replace("Error: ", "", 1).replace(f" {path} ", " it ")
        return None, why or "cannot be read"
    except Exception as exc:
        return None, f"{type(exc).__name__}: {one_line(exc)}"


# ---------------------------------------------------------------------------
# Sources
# ---------------------------------------------------------------------------

class NotChecked(Exception):
    """A source that is skipped with a note instead of an error."""


class Sequence(NamedTuple):
    words: list     # casefolded words of one quote, one text file, one whole PDF or another chapter
    pages: list     # for each word: its page (str) or None; for another chapter its line or paragraph
    label: str      # how a match is named: "Brandt_2019.md Q3", "Brandt_2019.pdf"
    kind: str       # "extract", "pdf", "text", or "chapter" for another chapter (--chapters)
    note: str = ""  # shown with the page, e.g. "secondary citation"; another chapter's "line" or "paragraph"
    first: bool = True  # False for the later parts of a quote split at an omission


WORDS = {}  # one str object per distinct word: a book-length source stays small


def source_tokens(text: str) -> list:
    text = PANDOC_CITATION.sub(" ", CITATION.sub(" ", text))
    return [WORDS.setdefault(w, w) for w in tokens(text)]


PAGE = re.compile(r"(?i)(?:\bpages?|\bseiten?|\bpp?\.|\bs\.)\s*((?:\d+|[ivxlc]+\b)"
                  r"(?:\s*(?:[-–]|to|bis)\s*(?:\d+|[ivxlc]+\b))?(?:\s*ff?\.)?)")


def page_of(text: str) -> Optional[str]:
    """The page named in a heading or bullet, with a ❓ the extract put there."""
    m = PAGE.search(text)
    if not m:
        return None
    page = re.sub(r"\s+", " ", m.group(1))
    return page + " ❓" if "❓" in text else page


FIELD = re.compile(r"\s*(?:[-*+]\s*)?\*\*([^*:]+?):?\*\*:?\s*(.*)")
# A line that starts a field ends the quote above it, with or without a bullet;
# the colon or "?" keeps a quote line in bold emphasis from counting as one.
LABEL = re.compile(r"\s*(?:[-*+]\s*)?\*\*[^*\n]+?(?::\*\*|\*\*\s*:|\?\*\*)")
# The bullet layout of the research-agent's rule examples: - p. 42: "quote"
PAGE_BULLET = re.compile(r"\s*[-*+]\s*((?:pp?\.|S\.|pages?|Seiten?)\s*[^:]{1,30}?)\s*:\s*(?=[\"„“»«'‘‚])(.*)", re.I)
# A page bullet without a quote, as in Rule 7: - p. 87 ❓ (e-book pagination)
PAGE_LINE = re.compile(r"\s*[-*+]\s*(?:pp?\.|S\.|pages?|Seiten?)\s*(?:\d|[ivxlc]+\b)", re.I)
# What makes an .md file an extract (or a quote registry) rather than a note
EXTRACT_MARK = re.compile(r"\*\*Original|^\s{0,3}#{1,6}\s+(?:Q\d+|Pages?\b|Seiten?\b)"
                          r"|^\s*\|.*(?:original|zitat|quote)", re.I | re.M)


def read_extract(path: Path) -> list:
    """
    The quotes of an extract, each with its page and quote ID: the
    "Original" field, quoted text in the "Secondary citation?" field, and
    "- p. 42: \"...\"" bullets. A Markdown table with an "Original quote"
    column (the source registry of workspace/state.md in seminar projects)
    is read row by row.
    """
    lines = normalize(read_file_text(path)).replace("\r\n", "\n").split("\n")
    seqs = []
    page, qid = None, None
    section_page = None  # the page of the last heading that is not a quote ID

    def add(text: str, where: Optional[str], label: str, note: str = "") -> None:
        # the parts of a quote between […] omissions are not one run in the source
        first = True
        for part in quote_parts(text):
            words = source_tokens(part)
            if words:
                seqs.append(Sequence(words, [where] * len(words), label, "extract", note, first))
                first = False

    def continuation(i: int, first: str) -> tuple:
        parts = [first]
        while (i < len(lines) and lines[i].strip() and not re.match(r"\s*(?:[-*+]\s+|#|\|)", lines[i])
               and not LABEL.match(lines[i])):
            parts.append(lines[i].strip())
            i += 1
        return " ".join(parts), i

    i = 0
    while i < len(lines):
        line = lines[i]
        h = re.match(r"\s{0,3}#{1,6}\s+(.*)", line)
        if h:
            heading = h.group(1).strip()
            q = re.match(r"(Q\d+)\b", heading)
            if q:
                # a page bullet inside the previous entry ends with it
                qid = q.group(1)
                page = page_of(heading[q.end():]) or section_page
            else:
                # any other heading starts a new page, or a part without one
                qid = None
                page = section_page = page_of(heading)
            i += 1
            continue
        if line.lstrip().startswith("|"):
            i = read_table(lines, i, path, add)
            continue
        label = path.name + (f" {qid}" if qid else "")
        f = FIELD.match(line)
        if f:
            name, rest = f.group(1).strip().lower(), f.group(2)
            if name == "original" or name.startswith(("original ", "original(")):
                text, i = continuation(i + 1, rest)
                add(text, page, label)
                continue
            elif name.startswith("secondary citation"):
                text, i = continuation(i + 1, rest)
                for m in QUOTED.finditer(text):
                    add(m.group(0)[1:-1], page, label, "secondary citation")
                continue
        b = PAGE_BULLET.match(line)
        if b:
            text, i = continuation(i + 1, b.group(2))
            add(text, page_of(b.group(1)) or page, label)
            continue
        if PAGE_LINE.match(line):
            page = page_of(line) or page
        i += 1
    return seqs


def looks_like_extract(path: Path) -> bool:
    return bool(EXTRACT_MARK.search(read_file_text(path)))


def read_table(lines: list, i: int, path: Path, add) -> int:
    """A Markdown table; rows of one with an "Original quote" column are quotes."""
    rows = []
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(rows) < 2:
        return i
    head = [c.lower() for c in rows[0]]
    quote = next((n for n, c in enumerate(head) if "original" in c or c in ("quote", "zitat")), None)
    if quote is None:
        return i
    page = next((n for n, c in enumerate(head) if c in ("page", "seite", "p.", "s.")), None)
    source = next((n for n, c in enumerate(head) if c in ("source", "quelle")), None)
    for row in rows[1:]:
        if len(row) <= quote or re.fullmatch(r"[\s:-]*", "".join(row)):
            continue  # the separator row |---|---|
        where = row[page] if page is not None and page < len(row) and row[page] else None
        name = row[source] if source is not None and source < len(row) and row[source] else ""
        add(row[quote], where, f"{path.name}: {name}" if name else path.name)
    return i


def line_key(line: str) -> str:
    """A line with its numbers masked: "Okafor 2021 · 117" and "... · 118" match."""
    return re.sub(r"\d+", "#", " ".join(line.split()).casefold())


def running_lines(texts: list) -> set:
    """
    The running headers and footers of a PDF: short lines (numbers masked)
    that open or close at least two pages and half of the left or of the right
    pages, since a book alternates two heads.
    """
    counts = {}
    for n, text in enumerate(texts):
        lines = [line for line in text.split("\n") if line.strip()]
        for key in {line_key(line) for line in lines[:2] + lines[-2:] if len(line.split()) <= 12}:
            counts.setdefault(key, [0, 0])[n % 2] += 1
    sides = [(len(texts) + 1) // 2, len(texts) // 2]
    return {key for key, c in counts.items()
            if sum(c) >= 2 and any(c[i] and c[i] * 2 >= sides[i] for i in (0, 1))}


def folio_head(line: str) -> bool:
    """A short line that starts or ends with a bare page number: "118  Kühle Höfe"."""
    words = line.split()
    return 2 <= len(words) <= 10 and bool(re.fullmatch(r"\d{1,4}", words[0]) or re.fullmatch(r"\d{1,4}", words[-1]))


def read_pdf(path: Path, vocabulary: set, unchecked: list) -> list:
    try:
        import logging
        from pypdf import PdfReader
        # pypdf's notes on fonts it cannot decode would land on stderr ahead of the report
        logging.getLogger("pypdf").setLevel(logging.CRITICAL + 1)
    except ImportError:
        raise NotChecked("pypdf not installed: pip install -r requirements.txt")
    with open(path, "rb") as fh:
        if fh.read(4) == b"PK\x03\x04":
            raise NotChecked("a ZIP archive named .pdf: extract it with scripts/extract_styleguide.sh "
                             "into a folder outside the sources")
    reader = PdfReader(str(path))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("encrypted with a password")
    def joins(a: str, b: str) -> bool:
        # A word hyphenated at a line end ("Verdun-" + "stung") is joined when the
        # chapter has the joined word; otherwise the hyphen is that of a compound
        # ("long-" + "term", "Stadt-" + "und") and the parts stay apart.
        return b[:1].islower() and (a + b).casefold() in vocabulary

    texts = [pg.extract_text() or "" for pg in reader.pages]
    running = running_lines(texts)
    words, pages = [], []
    hyphen = False  # the previous page ended in a hyphenated word
    empty = 0       # pages without any text: scanned (or blank) pages
    for n, raw in enumerate(texts, 1):
        # Running headers, footers and page numbers are dropped, so that a run
        # (and a word hyphenated at the page end) goes on across the page break.
        lines = [line for line in raw.replace("\r\n", "\n").split("\n") if line.strip()]
        for _ in range(3):
            if lines and (PAGE_NUMBER_LINE.fullmatch(lines[0]) or line_key(lines[0]) in running):
                lines.pop(0)
            if lines and (PAGE_NUMBER_LINE.fullmatch(lines[-1]) or line_key(lines[-1]) in running):
                lines.pop()
        # in a short excerpt a head appears too seldom to repeat; its page number gives it away
        if len(texts) <= 4:
            if len(lines) > 1 and folio_head(lines[0]):
                lines.pop(0)
            if len(lines) > 1 and folio_head(lines[-1]):
                lines.pop()
        # before normalize(), which drops soft hyphens (U+00AD)
        text = re.sub(r"([^\W\d_]+)[-\u00ad\u2010][ \t]*\n[ \t]*([^\W\d_]+)",
                      lambda m: m.group(1) + m.group(2) if joins(m.group(1), m.group(2)) else m.group(0),
                      "\n".join(lines))
        end = re.search(r"([^\W\d_]+)[-\u00ad\u2010]\s*$", text)
        text = normalize(text)
        page_words = source_tokens(text)
        empty += not WORD.search(raw)  # no text layer at all, not a page of only a head
        # ... and so is one hyphenated across a page break
        if hyphen and page_words and words and joins(words[-1], page_words[0]):
            joined = words[-1] + page_words.pop(0)
            words[-1] = WORDS.setdefault(joined, joined)
        hyphen = bool(end)
        words += page_words
        pages += [str(n)] * len(page_words)
    total = len(texts)
    if words and empty >= 2 and empty * 2 >= total:
        unchecked.append(f"{path} ({empty} of {total} pages without a text layer, scanned pages? "
                         "Those pages were not compared)")
    return [Sequence(words, pages, path.name, "pdf")] if words else []


def read_text(path: Path) -> list:
    words = source_tokens(normalize(read_file_text(path)))
    return [Sequence(words, [None] * len(words), path.name, "text")] if words else []


def source_files(arg: str, unchecked: list) -> list:
    """The source files of one argument; what cannot be read goes to unchecked."""
    p = Path(arg)
    if p.is_dir():
        found = []
        for root, dirs, files in os.walk(p):
            dirs[:] = sorted(d for d in dirs if d not in SKIPPED_FOLDERS and not d.startswith("."))
            # a saved web page's assets folder ("Article_files") holds no source
            dirs[:] = [d for d in dirs if not d.endswith("_files")]
            for f in sorted(files):
                if f.startswith((".", "~$")):
                    continue  # hidden files, macOS "._" forks, Office lock files
                suffix = Path(f).suffix.lower()
                if suffix in SOURCE_SUFFIXES:
                    found.append(Path(root) / f)
                elif suffix in LISTED_SUFFIXES:
                    unchecked.append(f"{Path(root) / f} ({suffix} is not read: give a PDF or text version)")
        return found
    try:
        with open(p, "rb"):
            pass
    except OSError as exc:
        reason = "not found" if isinstance(exc, FileNotFoundError) else exc.strerror or str(exc)
        unchecked.append(f"{p} ({reason})")
        return []
    if p.suffix.lower() not in SOURCE_SUFFIXES:
        unchecked.append(f"{p} ({p.suffix or 'a file without extension'} is not read: "
                         "use .md extracts, .pdf or .txt)")
        return []
    return [p]


# ---------------------------------------------------------------------------
# Matching
# ---------------------------------------------------------------------------

def find_runs(chapter: Chapter, seqs: list) -> list:
    """
    Maximal runs of at least POLISH words (tokens, counted as written later):
    [(start, length, [(seq, pos, length, start)])].
    """
    ch = chapter.words
    # The chapter's word windows are indexed, not the sources': a few thousand
    # entries instead of one per source word, however long the sources are.
    index = {}
    for i in range(len(ch) - POLISH + 1):
        window = ch[i:i + POLISH]
        if None not in window:
            index.setdefault(tuple(window), []).append(i)
    runs = {}
    for s, seq in enumerate(seqs):
        w = seq.words
        for p in range(len(w) - POLISH + 1):
            for i in index.get(tuple(w[p:p + POLISH]), ()):
                if i > 0 and p > 0 and ch[i - 1] is not None and ch[i - 1] == w[p - 1]:
                    continue  # part of a run that starts earlier
                n = POLISH
                while i + n < len(ch) and p + n < len(w) and ch[i + n] is not None and ch[i + n] == w[p + n]:
                    n += 1
                runs.setdefault((i, n), []).append((s, p))
    kept = []
    for (i, n), hits in sorted(runs.items(), key=lambda r: -r[0][1]):
        outer = next((k for k in kept if k[0] <= i and i + n <= k[0] + k[1]), None)
        if outer is None:
            kept.append((i, n, [(s, p, n, i) for s, p in hits]))
        else:
            # a shorter run inside a longer one, from another source, is named too:
            # an extract quote carries the quote ID and printed page a PDF lacks
            named = {h[0] for h in outer[2]}
            outer[2].extend((s, p, n, i) for s, p in hits if s not in named)
    return sorted(kept)


def written_ids(chapter: Chapter) -> list:
    """For each entry of chapter.words, the number of the word as written (the
    stretch between whitespace or dashes) it falls in, or None where a run
    breaks."""
    starts = [m.start() for m in re.finditer(r"\S+", spaced(chapter.text))]
    return [None if sp is None else bisect.bisect_right(starts, sp[0]) - 1 for sp in chapter.spans]


def run_length(chapter: Chapter, ids: list, i: int, n: int) -> tuple:
    """
    The length of the run chapter.words[i:i + n] in words as written: those its
    tokens fall in outside a short quotation, and those only inside one.
    """
    outside = {ids[k] for k in range(i, i + n) if not chapter.quoted[k]}
    inside = {ids[k] for k in range(i, i + n) if chapter.quoted[k]} - outside
    return len(outside), len(inside)


def quote_places(parts: list, seqs: list) -> dict:
    """For each part (a word list of POLISH or more words) found verbatim in a
    sequence, its index: [(seq, pos)] of the places it was found at."""
    starts = {}
    for k, words in enumerate(parts):
        starts.setdefault(tuple(words[:POLISH]), []).append(k)
    places = {}
    for s, seq in enumerate(seqs):
        w = seq.words
        for p in range(len(w) - POLISH + 1):
            for k in starts.get(tuple(w[p:p + POLISH]), ()):
                if w[p:p + len(parts[k])] == parts[k]:
                    places.setdefault(k, []).append((s, p))
    return places


def found_quotes(parts: list, seqs: list) -> set:
    """The indexes of the parts (word lists of POLISH or more words) found verbatim in a source."""
    return set(quote_places(parts, seqs))


def describe(seqs: list, hits: list, n: int, written) -> str:
    """The sources of a run of n tokens; written(start, m): m tokens as written."""
    names = []
    for s, p, m, i in sorted(hits, key=lambda h: (h[2] < n, seqs[h[0]].kind != "extract")):
        seq = seqs[s]
        pages = [pg for pg in seq.pages[p:p + m] if pg]
        if seq.kind == "pdf" and pages:
            where = [f"PDF page {pages[0]}" if pages[0] == pages[-1] else f"PDF pages {pages[0]}-{pages[-1]}"]
        else:
            where = [f"p. {pages[0]}"] if pages else []
        where += [seq.note] if seq.note else []
        where += [f"{written(i, m)} of these words"] if m < n else []
        name = f"{seq.label} ({', '.join(where)})" if where else seq.label
        if name not in names:
            names.append(name)
    shown = "; ".join(names[:3])
    return shown + (f"; and {len(names) - 3} more" if len(names) > 3 else "")


def describe_chapters(seqs: list, hits: list, n: int, written) -> str:
    """
    Where a run of n tokens stands in the other chapters: "ch_2_basics.md
    (line 18)", "ch_2_basics.md (lines 18, 40)"; a chapter that has only part
    of it, with the number of these words, as describe() names a source.
    """
    places = {}  # (file name, unit, words of a part or None) -> [line or paragraph]
    for s, p, m, i in sorted(hits, key=lambda h: (h[2] < n, h[0], h[1])):
        seq = seqs[s]
        where = places.setdefault((seq.label, seq.note, written(i, m) if m < n else None), [])
        if seq.pages[p] not in where:
            where.append(seq.pages[p])
    names = []
    for (label, unit, part), where in places.items():
        at = f"{unit}s {', '.join(where)}" if len(where) > 1 else f"{unit} {where[0]}"
        names.append(f"{label} ({at}" + (f", {part} of these words)" if part else ")"))
    shown = "; ".join(names[:3])
    return shown + (f"; and {len(names) - 3} more" if len(names) > 3 else "")


def shorten(text: str) -> str:
    """One line; a long text keeps its first and last words around " … "."""
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r" ([.,;:!?])", r"\1", text)  # a blanked citekey leaves "word ."
    if len(text) <= 160:
        return text
    return text[:120].rsplit(" ", 1)[0] + " … " + text[-40:].split(" ", 1)[-1]


def rating(n: int) -> str:
    return "❌" if n >= SERIOUS else "⚠️ " if n >= RISK else "🟡"


def main() -> None:
    args = sys.argv[1:]
    if any(a in ("-h", "--help") for a in args):
        print((__doc__ or USAGE).strip())
        sys.exit(0)
    if "--words" in args:
        # the word count alone, for the page budget: no sources needed
        rest = [a for a in args if a != "--words"]
        if len(rest) != 1 or args.count("--words") != 1:
            print(USAGE, file=sys.stderr)
            sys.exit(2)
        print(word_count(read_chapter(checked_chapter(Path(rest[0])))))
        sys.exit(0)
    # the other chapters, for A.2 of checklists/consistency.md: the paths after --chapters
    other_args = None
    if "--chapters" in args:
        at = args.index("--chapters")
        args, other_args = args[:at], args[at + 1:]
        if "--chapters" in other_args:
            print(USAGE, file=sys.stderr)
            sys.exit(2)
    if len(args) < 2:
        print(USAGE, file=sys.stderr)
        sys.exit(2)

    chapter_path = checked_chapter(Path(args[0]))
    chapter = read_chapter(chapter_path)
    if not any(w is not None for w in chapter.words) and not chapter.quotes:
        fail(f"no text found in {chapter_path}")
    # the chapter's words, quotations included, decide hyphen joins in the PDFs
    vocabulary = {w for w in chapter.words if w is not None}
    vocabulary.update(w for _, q in chapter.quotes for w in tokens(q))

    files, unchecked = [], []
    named = set()  # files given as arguments, not found in a folder
    for arg in args[1:]:
        found = [f for f in source_files(arg, unchecked) if f.resolve() != chapter_path.resolve()]
        if not Path(arg).is_dir():
            named.update(f.resolve() for f in found)
        files += found

    seqs, unreadable = [], []
    counts = {"extract": 0, "pdf": 0, "text": 0}
    for f in files:
        suffix = f.suffix.lower()
        try:
            if suffix == ".pdf":
                found = read_pdf(f, vocabulary, unchecked)
                if not found:
                    unchecked.append(f"{f} (no text layer, a scan?)")
            elif suffix == ".txt":
                found = read_text(f)
            else:
                found = read_extract(f)
                # a note found in a folder (handover.md, README.md) is not a source
                if not found and (f.resolve() in named or looks_like_extract(f)):
                    unchecked.append(f"{f} (no quotes found: not an extract?)")
        except SystemExit:
            raise
        except NotChecked as note:
            unchecked.append(f"{f} ({note})")
            continue
        except Exception as exc:
            unreadable.append(f"{f} ({type(exc).__name__}: {one_line(exc)})")
            continue
        if found:
            counts[found[0].kind] += 1
        seqs += found

    if not seqs:
        # A clean result would claim a check that did not happen.
        fail("nothing to compare against: no extract quotes, PDF text or text files in "
             + ", ".join(args[1:])
             + "".join(f"; {u}" for u in unchecked + unreadable))
    # A run is rated by its words as written (a hyphenated compound is one word,
    # a dash between words separates them) outside quotation marks: the words of
    # a short quotation inside it keep it going but are not counted.
    ids = written_ids(chapter)
    runs = []
    for i, n, hits in find_runs(chapter, seqs):
        c, q = run_length(chapter, ids, i, n)
        if c >= POLISH:
            runs.append((i, n, c, q, hits))

    # The other chapters (--chapters), read as the chapter is: a run they share
    # with it is counted the same way, and their quotations, block quotes,
    # bibliography and raw blocks take no part. What cannot be read goes into
    # the lists of the sources' report, but not into its line on the sources.
    others, other_unchecked, other_unreadable = [], [], []
    for f in other_chapter_files(other_args or [], chapter_path, other_unchecked):
        if f.suffix.lower() == ".docx":
            # without python-docx a Word chapter is not checked, as a PDF source is without pypdf
            try:
                import docx  # noqa: F401
            except ImportError:
                other_unchecked.append(f"{f} (python-docx not installed: pip install -r requirements.txt)")
                continue
        other, why = read_other_chapter(f)
        if other is None:
            other_unreadable.append(f"{f} ({why})")
        else:
            others.append(Sequence(other.words, [None if pl is None else str(pl) for pl in other.places],
                                   f.name, "chapter", other.unit))
    # A run that lies wholly within runs listed against the sources is not
    # listed again: step 8 of the source-verifier rates it.
    from_sources = {k for i, n, *_ in runs for k in range(i, i + n)}
    shared = []
    for i, n, hits in find_runs(chapter, others):
        c, q = run_length(chapter, ids, i, n)
        if c >= POLISH and not all(k in from_sources for k in range(i, i + n)):
            also = len({ids[k] for k in range(i, i + n) if k in from_sources and not chapter.quoted[k]})
            shared.append((i, n, c, q, also, hits))

    print(f"Verbatim check for: {chapter_path.name}")
    print(f"  Sources: {counts['extract']} extract(s) with "
          f"{sum(1 for s in seqs if s.kind == 'extract' and s.first)} quotes, "
          f"{counts['pdf']} PDF(s), {counts['text']} text file(s)")
    if other_args is not None:
        print(f"  Other chapters: {len(others)}" + (f" ({', '.join(s.label for s in others)})" if others else ""))
    print(f"  {word_count(chapter)}")
    if chapter.bibliography:
        first, last = chapter.bibliography
        print(f"  Not counted: the bibliography, " + (f"{chapter.unit}s {first}-{last}" if last
                                                      else f"from {chapter.unit} {first} on"))
    print()

    for i, n, c, q, hits in runs:
        start, end = chapter.spans[i][0], chapter.spans[i + n - 1][1]
        excerpt = shorten(chapter.text[start:end])
        also = f" (and {q} in quotation marks)" if q else ""
        print(f"{rating(c)} {c} words{also}, {chapter.unit} {chapter.places[i]}: \"{excerpt}\"")
        print(f"    in {describe(seqs, hits, n, lambda i, m: sum(run_length(chapter, ids, i, m)))}")
    if runs:
        print()

    checks = []  # (place, quote, words of one part between omissions)
    for place, quote in chapter.quotes:
        for part in quote_parts(quote):
            words = tokens(PANDOC_CITATION.sub(" ", CITATION.sub(" ", part)))
            if len(words) >= POLISH:
                checks.append((place, quote, words))
    found = found_quotes([words for _, _, words in checks], seqs)
    missing = []
    for k, (place, quote, _) in enumerate(checks):
        if k not in found and (place, quote) not in missing:
            missing.append((place, quote))
    if missing:
        print("Direct quotes not found verbatim in the sources (check by hand; "
              "the source may be missing here):")
        for place, quote in missing:
            print(f"    {chapter.unit} {place}: \"{shorten(quote)}\"")
        print()

    if unchecked or other_unchecked:
        print("Not checked:")
        for u in unchecked + other_unchecked:
            print(f"    {u}")
        print()
    if unreadable or other_unreadable:
        print("Could not be read:")
        for u in unreadable + other_unreadable:
            print(f"    {u}")
        print()

    serious = sum(1 for run in runs if run[2] >= SERIOUS)
    risk = sum(1 for run in runs if RISK <= run[2] < SERIOUS)
    polish = sum(1 for run in runs if run[2] < RISK)
    if runs:
        print(f"Summary: ❌ {serious}, ⚠️ {risk}, 🟡 {polish}. Runs of 1-4 words are not listed;")
    else:
        print("✅ No verbatim run of 5 or more words outside quotation marks"
              + (" in the sources that were read." if unchecked or unreadable else "."))
    print("text in quotation marks and block quotes is not counted (see checklists/plagiarism_protection.md).")

    if other_args is not None:
        # After the part on the sources, which stays as it is without --chapters.
        # A direct quote that no source has but another chapter has as its own
        # text is a quotation of that chapter (the research question quoted again).
        in_others = quote_places([words for _, _, words in checks], others)
        quoted = []  # (place, quote, [(seq, pos)])
        for k, (place, quote, _) in enumerate(checks):
            if k in found or k not in in_others:
                continue
            if quoted and quoted[-1][:2] == (place, quote):
                quoted[-1][2].extend(in_others[k])
            else:
                quoted.append((place, quote, list(in_others[k])))
        print()
        if not shared and not quoted:
            print("Runs shared with other chapters (A.2 of checklists/consistency.md): "
                  "none of 5 or more words, and no quotation of another chapter's text.")
        else:
            print("Runs shared with other chapters (A.2 of checklists/consistency.md; not rated):")
        for i, n, c, q, also, hits in shared:
            start, end = chapter.spans[i][0], chapter.spans[i + n - 1][1]
            inside = f" (and {q} in quotation marks)" if q else ""
            print(f"{c} words{inside}, {chapter.unit} {chapter.places[i]}: \"{shorten(chapter.text[start:end])}\"")
            print(f"    in {describe_chapters(others, hits, n, lambda i, m: sum(run_length(chapter, ids, i, m)))}"
                  + (f"; {also} of these words also in a run from the sources above" if also else ""))
        for place, quote, at in quoted:
            print(f"quotation, {chapter.unit} {place}: \"{shorten(quote)}\"")
            # each place has the whole quotation part (n = m = 0: no "of these words")
            print(f"    in {describe_chapters(others, [(s, p, 0, 0) for s, p in at], 0, None)}, "
                  "there outside quotation marks")

    if unreadable or other_unreadable:
        sys.exit(2)
    sys.exit(1 if serious or risk else 0)


if __name__ == "__main__":
    main()
